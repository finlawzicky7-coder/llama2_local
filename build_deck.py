#!/usr/bin/env python3
"""
Generate a downloadable, editable PowerPoint investor pitch deck for
Ridgewell Insurance Agency LLC — a Medicare agency whose core moat is
AI-driven Renewal & Retention Intelligence.

Output: ridgewell-investor-pitch.pptx (16:9)

Run: python3 build_deck.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
NAVY   = RGBColor(0x0E, 0x2A, 0x3B)   # deep slate-navy
BLUE   = RGBColor(0x1B, 0x6E, 0xC2)   # accent blue
GREEN  = RGBColor(0x1E, 0x7A, 0x4B)   # ridge / evergreen
TEAL   = RGBColor(0x14, 0xA3, 0x8B)   # supporting teal
AMBER  = RGBColor(0xF2, 0x9F, 0x05)   # warm accent
INK    = RGBColor(0x1A, 0x1A, 0x1A)
SLATE  = RGBColor(0x5B, 0x6B, 0x7B)
LIGHT  = RGBColor(0xF4, 0xF7, 0xFA)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT= RGBColor(0xEA, 0xF1, 0xF8)

FONT = "Calibri"
BRAND = "Ridgewell Insurance Agency LLC  ·  Renewal & Retention Intelligence"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
gap = Inches(0.25)


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    return tb, tf


def setpara(p, text, size, color, bold=False, align=PP_ALIGN.LEFT,
            space_after=6, italic=False, font=FONT):
    p.text = text if text else " "
    p.alignment = align; p.space_after = Pt(space_after)
    r = p.runs[0]
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    return p


def add_para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT,
             space_after=6, italic=False, bullet=False, level=0):
    p = tf.add_paragraph(); p.level = level
    setpara(p, ("• " + text) if bullet else text, size, color, bold,
            align, space_after, italic)
    return p


def header(s, kicker, title, num):
    rect(s, 0, 0, Inches(0.28), SH, BLUE)
    tb, tf = textbox(s, Inches(0.6), Inches(0.35), Inches(11.8), Inches(1.3))
    setpara(tf.paragraphs[0], kicker, 13, GREEN, bold=True, space_after=2)
    add_para(tf, title, 29, NAVY, bold=True, space_after=0)
    nb, nf = textbox(s, Inches(12.4), Inches(6.95), Inches(0.8), Inches(0.4))
    setpara(nf.paragraphs[0], str(num), 11, SLATE, align=PP_ALIGN.RIGHT)
    fb, ff = textbox(s, Inches(0.6), Inches(6.95), Inches(10), Inches(0.4))
    setpara(ff.paragraphs[0], BRAND, 10, SLATE)


def table(s, rows, cols, x, y, w, h, data, header_fill=NAVY,
          col_widths=None, font_size=12, header_size=12):
    gt = s.shapes.add_table(rows, cols, x, y, w, h).table
    if col_widths:
        for i, cw in enumerate(col_widths):
            gt.columns[i].width = cw
    for r in range(rows):
        for c in range(cols):
            cell = gt.cell(r, c)
            cell.margin_left = Inches(0.1); cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            is_head = (r == 0)
            setpara(tf.paragraphs[0], data[r][c],
                    header_size if is_head else font_size,
                    WHITE if is_head else INK, bold=is_head, space_after=0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (header_fill if is_head
                                        else (ROW_ALT if r % 2 else WHITE))
    return gt


def card(s, x, y, w, h, title, body_lines, accent=BLUE, title_size=16):
    rect(s, x, y, w, h, LIGHT)
    rect(s, x, y, w, Inches(0.12), accent)
    tb, tf = textbox(s, x + Inches(0.2), y + Inches(0.28),
                     w - Inches(0.4), h - Inches(0.4))
    setpara(tf.paragraphs[0], title, title_size, NAVY, bold=True, space_after=6)
    for ln in body_lines:
        add_para(tf, ln, 12.5, INK, space_after=5)
    return tb


# ================================================================ SLIDE 1
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(5.0), SW, Inches(0.10), AMBER)
rect(s, Inches(0.9), Inches(1.55), Inches(0.16), Inches(2.0), GREEN)
tb, tf = textbox(s, Inches(1.25), Inches(1.45), Inches(11.3), Inches(3.4))
setpara(tf.paragraphs[0], "Ridgewell Insurance Agency", 50, WHITE, bold=True, space_after=0)
add_para(tf, "Renewal & Retention Intelligence for Medicare.", 27, AMBER,
         bold=True, space_after=18)
add_para(tf, "We sell Medicare — and we keep members enrolled, using AI-driven",
         19, WHITE, space_after=2)
add_para(tf, "retention intelligence that compounds renewal commissions.", 19, WHITE,
         space_after=0)
tb2, tf2 = textbox(s, Inches(1.25), Inches(6.4), Inches(11.3), Inches(0.8))
setpara(tf2.paragraphs[0],
        "Seed / Series A  ·  Health-Tech Investor Pitch  ·  Confidential",
        14, RGBColor(0xC9, 0xD8, 0xE8))

# ================================================================ SLIDE 2  Problem
s = slide()
header(s, "THE PROBLEM", "In Medicare, the sale is just the beginning — retention is the business", 2)
cw = Inches(3.85); x0 = Inches(0.6); y0 = Inches(2.0); ch = Inches(3.7)
card(s, x0, y0, cw, ch, "Churn destroys commission value",
     ["Rapid disenrollment claws back the agency's commission entirely.",
      "Every member who switches at AEP resets lifetime value to zero.",
      "Renewal commissions only accrue while the member stays."], BLUE)
card(s, x0 + cw + gap, y0, cw, ch, "Renewal season is chaos",
     ["Each AEP/OEP, ANOC changes, plan exits, and benefit shifts confuse members.",
      "Confusion — not dissatisfaction — drives most plan switching.",
      "Members rarely understand what actually changed in their plan."], TEAL)
card(s, x0 + 2 * (cw + gap), y0, cw, ch, "Agencies fly blind",
     ["Most agencies have no early-warning system for at-risk members.",
      "Outreach is reactive, manual, and badly timed.",
      "Retention is left to luck instead of intelligence."], AMBER)
tb, tf = textbox(s, Inches(0.6), Inches(5.95), Inches(12), Inches(0.8))
setpara(tf.paragraphs[0],
        "A few points of persistency is the difference between a shrinking book and a compounding one.",
        15, NAVY, bold=True, italic=True)

# ================================================================ SLIDE 3  Why Now
s = slide()
header(s, "WHY NOW", "Policy, market, and technology favor a retention-led agency", 3)
cw = Inches(3.85); x0 = Inches(0.6); y0 = Inches(2.0); ch = Inches(3.9)
card(s, x0, y0, cw, ch, "Policy tailwind",
     ["CMS is tightening MA marketing rules and scrutinizing rapid disenrollment.",
      "Pressure is shifting value from churn-and-burn selling to durable retention.",
      "Health-equity & language-access priorities reward clear member communication."], BLUE)
card(s, x0 + cw + gap, y0, cw, ch, "Market tailwind",
     ["MA covers ~half of all Medicare beneficiaries and keeps growing.",
      "~10,000 Americans age into Medicare every day.",
      "33M+ MA members face a renewal decision every single year."], TEAL)
card(s, x0 + 2 * (cw + gap), y0, cw, ch, "Technology tailwind",
     ["AI can now predict at-risk members from plan and behavior data.",
      "Plain-language translation + real-time voice are production-ready.",
      "Grounded, cited answers make member guidance safe and auditable."], AMBER)

# ================================================================ SLIDE 4  Solution
s = slide()
header(s, "THE SOLUTION", "Renewal & Retention Intelligence — our core engine", 4)
cw = Inches(5.85); x0 = Inches(0.6); y0 = Inches(2.0); ch = Inches(4.3)
card(s, x0, y0, cw, ch, "1 · Retention intelligence",
     ["Predictive at-risk scoring across the entire book of business.",
      "Flags members likely to disenroll or switch — before AEP, not after.",
      "Tracks each member's ANOC / plan / network / drug changes automatically.",
      "Prioritizes and times agent outreach for maximum retention."], GREEN, 18)
card(s, x0 + cw + gap, y0, cw, ch, "2 · Member navigation AI",
     ["Plain-language translation of plan documents at a ≤6th-grade level.",
      "Real-time conversational AI — by text or voice, in the member's language.",
      "Grounded in the member's actual plan, with citations to the source.",
      "Turns renewal confusion into a confident decision to stay."], TEAL, 18)

# ================================================================ SLIDE 5  Demo
s = slide()
header(s, "HOW IT SAVES A MEMBER", "A retention save, end to end", 5)
rect(s, Inches(0.6), Inches(1.9), Inches(7.2), Inches(4.6), NAVY)
tb, tf = textbox(s, Inches(0.95), Inches(2.15), Inches(6.6), Inches(4.1))
setpara(tf.paragraphs[0], "⚠️  At-risk flag (pre-AEP):", 15, AMBER, bold=True, space_after=4)
add_para(tf, "“Mrs. Alvarez's 2025 plan is dropping her cardiologist from the network.”",
         16, WHITE, space_after=12)
add_para(tf, "Agent outreach, AI-assisted (Spanish):", 13, RGBColor(0xC9,0xD8,0xE8), bold=True, space_after=4)
add_para(tf, "“Su médico sale de la red en enero. Encontré 2 planes que lo mantienen "
             "por un costo similar. ¿Quiere que la inscriba?”",
         15, WHITE, space_after=10)
add_para(tf, "📎 Fuente: 2025 ANOC + directorio de proveedores",
         12, TEAL, italic=True, space_after=12)
add_para(tf, "✓ Member retained on a suitable plan — renewal commission preserved.",
         14, AMBER, bold=True, space_after=0)
card(s, Inches(8.1), Inches(1.9), Inches(4.6), Inches(4.6),
     "Why this wins",
     ["Caught the risk before the member silently lapsed.",
      "Plain-language, multilingual explanation built trust.",
      "Agent kept the household — and the residual commission.",
      "Every step logged for CMS marketing-rule compliance.",
      "",
      "Reactive agencies never see this member until they're gone."], GREEN, 16)

# ================================================================ SLIDE 6  How it works
s = slide()
header(s, "HOW IT WORKS", "From book data to retained members", 6)
steps = ["Book + plan /\nANOC data", "At-risk\nscoring model",
         "Plain-language +\nconversational AI", "Guided outreach +\ncompliance layer",
         "Retained &\nrenewed member"]
n = len(steps); bw = Inches(2.18); bh = Inches(1.5); gap2 = Inches(0.24)
x = Inches(0.6); y = Inches(2.3)
for i, st in enumerate(steps):
    color = AMBER if i == 3 else (GREEN if i in (1,) else BLUE)
    r = rect(s, x, y, bw, bh, color)
    tf = r.text_frame; tf.word_wrap = True
    setpara(tf.paragraphs[0], st, 13, WHITE, bold=True, align=PP_ALIGN.CENTER)
    if i < n - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                x + bw, y + Inches(0.55), gap2, Inches(0.4))
        ar.fill.solid(); ar.fill.fore_color.rgb = SLATE
        ar.line.fill.background(); ar.shadow.inherit = False
    x = x + bw + gap2
card(s, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.1),
     "Built-in guardrails (the compliance layer)",
     ["Grounding + citations: every member answer traces to official plan documents.   "
      "•  Scope limiter: licensed agents advise; AI educates and navigates.",
      "CMS marketing-rule safe: neutral, accurate, documented.   "
      "•  PHI minimization + immutable audit logs.   •  Human-in-the-loop on every enrollment."],
     TEAL, 16)

# ================================================================ SLIDE 7  Market
s = slide()
header(s, "MARKET SIZE", "A growing book — and a platform other agencies need", 7)
cx = Inches(3.4); cy = Inches(4.3)
for r_in, color in [(2.4, RGBColor(0xD6,0xE6,0xF5)),
                    (1.7, RGBColor(0x9F,0xC4,0xE8)),
                    (1.0, GREEN)]:
    d = Inches(r_in*2)
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, cx-Inches(r_in), cy-Inches(r_in), d, d)
    c.fill.solid(); c.fill.fore_color.rgb = color
    c.line.fill.background(); c.shadow.inherit = False
def lbl(y, t, sz, col):
    b,f = textbox(s, Inches(1.3), y, Inches(4.2), Inches(0.6), MSO_ANCHOR.MIDDLE)
    setpara(f.paragraphs[0], t, sz, col, bold=True, align=PP_ALIGN.CENTER)
lbl(Inches(2.2), "TAM · ~68M beneficiaries; multibillion commission pool", 12, NAVY)
lbl(Inches(3.6), "SAM · 33M+ MA members renewing yearly + their agencies", 11.5, NAVY)
lbl(Inches(4.5), "SOM · Ridgewell's book + licensed agencies", 12, WHITE)
card(s, Inches(6.9), Inches(2.0), Inches(5.8), Inches(4.5),
     "Two ways to grow",
     ["1 · Grow our own book — every retained member compounds renewal commissions.",
      "2 · License the intelligence — sell Renewal & Retention Intelligence to other agencies & FMOs that face the same churn.",
      "",
      "Note: replace every figure with a current, sourced number (CMS / KFF / your book) before pitching."],
     TEAL, 16)

# ================================================================ SLIDE 8  Business model
s = slide()
header(s, "BUSINESS MODEL", "Recurring commissions, compounded by retention", 8)
cw = Inches(3.85); x0 = Inches(0.6); y0 = Inches(2.0); ch = Inches(3.5)
card(s, x0, y0, cw, ch, "Sell Medicare",
     ["CMS-regulated initial + renewal commissions on MA, Med Supp, and Part D.",
      "Renewal commissions recur every year a member stays."], BLUE)
card(s, x0+cw+gap, y0, cw, ch, "Retain members  (the moat)",
     ["Retention intelligence lifts persistency.",
      "Higher persistency compounds residual commissions and book value."], GREEN)
card(s, x0+2*(cw+gap), y0, cw, ch, "License the intelligence",
     ["Per-life or per-seat SaaS to partner agencies & FMOs.",
      "Recurring revenue beyond our own book."], AMBER)
tb, tf = textbox(s, Inches(0.6), Inches(5.7), Inches(12), Inches(1.0))
setpara(tf.paragraphs[0],
        "An MA member's value is an annuity — retention is what turns each sale into years of recurring commission.",
        15, NAVY, bold=True, italic=True)

# ================================================================ SLIDE 9  ROI
s = slide()
header(s, "UNIT ECONOMICS", "Persistency is the highest-ROI lever in the agency", 9)
data = [
    ["Value lever", "Mechanism", "Illustrative impact"],
    ["Persistency lift", "Keep members past chargeback window & through AEP",
     "Each +1 pt persistency = residual commissions retained"],
    ["Chargeback avoidance", "Catch at-risk members before rapid disenrollment",
     "Avoid full commission clawback per lapsed member"],
    ["Agent productivity", "Prioritized at-risk lists + AI member navigation",
     "More renewals per agent; less manual triage at AEP"],
    ["Cross-sell", "Trusted touchpoints surface Med Supp / PDP / ancillary",
     "Added commission per retained household"],
    ["Licensing", "Sell the intelligence to partner agencies",
     "Recurring SaaS revenue beyond our own book"],
]
table(s, 6, 3, Inches(0.6), Inches(2.0), Inches(12.1), Inches(3.0), data,
      col_widths=[Inches(2.6), Inches(4.6), Inches(4.9)], font_size=12)
card(s, Inches(0.6), Inches(5.2), Inches(12.1), Inches(1.5),
     "Lifetime-value math (CFO-editable)",
     ["Book LTV = members × (initial + Σ renewal commissions over retained years) × persistency. "
      "Retention multiplies every term that matters.",
      "Retaining a member 1–2 extra years compounds residual commissions far beyond acquisition cost — "
      "lifting persistency a few points across the book is the single biggest value driver."],
     TEAL, 15)

# ================================================================ SLIDE 10 Compliance
s = slide()
header(s, "REGULATORY & COMPLIANCE", "A licensed agency that treats compliance as a moat", 10)
phases = [
    ("Licensing & foundation", "State DOI agency/producer licenses, CMS marketing-rule "
     "compliant communications, HIPAA program + BAAs, encryption, audit logging", AMBER, "now"),
    ("Enterprise-ready", "SOC 2 Type I, WCAG 2.x AA + VPAT, penetration test, "
     "AI guardrails (educate not advise), incident-response runbook", BLUE, "before licensing to partners"),
    ("Scale", "SOC 2 Type II, AI governance + bias/accuracy monitoring, DPIAs, "
     "multi-state licensing expansion", TEAL, "Series A scale-up"),
    ("Moat", "Continuous compliance automation, third-party audits, "
     "CMS-aligned transparency reporting", GREEN, "Series A → B"),
]
y = Inches(2.1); rh = Inches(1.05)
for title, body, color, when in phases:
    rect(s, Inches(0.6), y, Inches(0.18), rh - Inches(0.15), color)
    rect(s, Inches(0.9), y, Inches(11.8), rh - Inches(0.15), LIGHT)
    tb, tf = textbox(s, Inches(1.1), y + Inches(0.05), Inches(11.4), rh - Inches(0.2))
    p = tf.paragraphs[0]
    setpara(p, title, 15, NAVY, bold=True, space_after=2)
    r2 = p.add_run(); r2.text = "    (" + when + ")"
    r2.font.size = Pt(12); r2.font.italic = True; r2.font.color.rgb = SLATE; r2.font.name = FONT
    add_para(tf, body, 12, INK, space_after=0)
    y = y + rh

# ================================================================ SLIDE 11 Traction
s = slide()
header(s, "TRACTION & MARKET READINESS", "A real book, an intelligence edge, ready to scale", 11)
cards = [
    ("Agency readiness", ["Licensed, producing Medicare agency with an active book.",
      "Carrier appointments in place.", "Renewal commissions already recurring."], BLUE),
    ("Intelligence readiness", ["At-risk scoring live on the book.",
      "Plain-language readability: grade 12.4 → 5.8.", "Persistency lift measured vs. baseline."], GREEN),
    ("Expansion readiness", ["Partner-agency demand to license the platform.",
      "LOIs / pilots with other agencies & FMOs.", "Multilingual member navigation deployed."], AMBER),
]
cw = Inches(3.85); x0 = Inches(0.6); y0 = Inches(2.0); ch = Inches(2.7)
for i,(t,lines,c) in enumerate(cards):
    card(s, x0 + i*(cw+gap), y0, cw, ch, t, lines, c)
rect(s, Inches(0.6), Inches(5.0), Inches(12.1), Inches(1.4), NAVY)
mets = [("Persistency", "vs. industry baseline"), ("Book size", "members / premium"),
        ("Retention lift", "from intelligence"), ("Languages", "supported")]
mw = Inches(12.1)/4
for i,(big,small) in enumerate(mets):
    bx = Inches(0.6) + mw*i
    tb,tf = textbox(s, bx, Inches(5.15), mw, Inches(1.1), MSO_ANCHOR.MIDDLE)
    setpara(tf.paragraphs[0], big, 18, AMBER, bold=True, align=PP_ALIGN.CENTER, space_after=2)
    add_para(tf, small, 12, WHITE, align=PP_ALIGN.CENTER, space_after=0)

# ================================================================ SLIDE 12 Competition
s = slide()
header(s, "COMPETITIVE LANDSCAPE", "Retention intelligence + compliance = defensible", 12)
data = [
    ["", "Predictive\nretention", "Member\nnavigation AI", "Multilingual\n≤6th grade",
     "Licensed\nagency", "Compliance-\nfirst"],
    ["Ridgewell", "✓", "✓", "✓", "✓", "✓"],
    ["Typical Medicare agencies", "—", "—", "partial", "✓", "partial"],
    ["CRM / lead-gen tools", "partial", "—", "—", "—", "partial"],
    ["FMO back-office platforms", "partial", "—", "partial", "—", "✓"],
    ["Generic LLM chatbots", "—", "partial", "✓", "—", "—"],
]
table(s, 6, 6, Inches(0.6), Inches(2.1), Inches(12.1), Inches(3.4), data,
      col_widths=[Inches(3.6)] + [Inches(1.7)]*5, font_size=12, header_size=11)
tb, tf = textbox(s, Inches(0.6), Inches(5.8), Inches(12), Inches(0.9))
setpara(tf.paragraphs[0],
        "Defensibility = proprietary book data + retention models + multilingual navigation + a licensed, compliant agency.",
        14, NAVY, bold=True, italic=True)

# ================================================================ SLIDE 13 Team
s = slide()
header(s, "TEAM", "The trifecta a retention-led Medicare agency needs", 13)
roles = [
    ("Medicare agency leadership", ["Builds the book, holds carrier appointments, knows AEP/OEP and persistency."], BLUE),
    ("Data / AI engineering", ["Builds at-risk models + grounded, multilingual member navigation."], GREEN),
    ("Regulatory / compliance", ["State licensing, CMS marketing rules, HIPAA, SOC 2."], AMBER),
]
cw = Inches(3.85); x0 = Inches(0.6); y0 = Inches(2.0); ch = Inches(2.8)
for i,(t,lines,c) in enumerate(roles):
    card(s, x0 + i*(cw+gap), y0, cw, ch, t, lines, c)
card(s, Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.4),
     "Advisors", ["Medicare distribution / FMO leaders, a data-science advisor, and healthcare "
                  "regulatory counsel — replace with your real names & logos."], NAVY, 16)

# ================================================================ SLIDE 14 Ask
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(2.4), SW, Inches(0.10), AMBER)
tb, tf = textbox(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(1.6))
setpara(tf.paragraphs[0], "The Ask", 44, WHITE, bold=True, space_after=4)
add_para(tf, "Raising $[X]M to grow the book and scale Renewal & Retention Intelligence.",
         20, AMBER, space_after=0)
cols = [("Grow the book", "members + premium"), ("Lift persistency", "vs. industry"),
        ("Launch licensing", "to partner agencies"), ("$[X] ARR", "in 18 months")]
mw = Inches(2.85); x = Inches(0.9); y = Inches(3.0)
for big, small in cols:
    rect(s, x, y, mw, Inches(1.5), RGBColor(0x16,0x3A,0x4E))
    tbx, tfx = textbox(s, x, y+Inches(0.2), mw, Inches(1.1), MSO_ANCHOR.MIDDLE)
    setpara(tfx.paragraphs[0], big, 17, WHITE, bold=True, align=PP_ALIGN.CENTER, space_after=2)
    add_para(tfx, small, 13, RGBColor(0xC9,0xD8,0xE8), align=PP_ALIGN.CENTER, space_after=0)
    x = x + mw + Inches(0.13)
tb, tf = textbox(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(2.0))
setpara(tf.paragraphs[0], "Use of funds", 18, AMBER, bold=True, space_after=6)
add_para(tf, "Agent growth & marketing  ·  Retention-intelligence engineering  ·  Compliance & licensing",
         16, WHITE, space_after=14)
add_para(tf, "Every dollar tied to a de-risking milestone the next round will underwrite.",
         15, RGBColor(0xC9,0xD8,0xE8), italic=True, space_after=0)
tb, tf = textbox(s, Inches(0.9), Inches(6.8), Inches(11.5), Inches(0.5))
setpara(tf.paragraphs[0], "Ridgewell Insurance Agency LLC  ·  finlawzicky7@gmail.com", 13, WHITE)

prs.save("ridgewell-investor-pitch.pptx")
print("Saved ridgewell-investor-pitch.pptx with", len(prs.slides._sldIdLst), "slides")
